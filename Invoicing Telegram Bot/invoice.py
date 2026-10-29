"""Turn an original CAMMED LAB invoice PDF into the 2-page Physiocam version.

Page 1 is the original page, untouched.
Page 2 is the same page with the "Total:" row replaced by
Original Amount / Total for customer / Discount / Total, and everything
below the table moved down to make room.

Both invoice formats that pdftext.py reads work: CAMMED's mPDF files and
invoices printed to PDF from a browser with the Adobe PDF printer.
"""

import io
import logging
import os
import re
import sys
from dataclasses import dataclass
from decimal import ROUND_HALF_EVEN, Decimal, InvalidOperation
from functools import cache
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from pdftext import UNKNOWN, Rule, Run, page_rules, page_runs

log = logging.getLogger(__name__)

DISCOUNT_RATE = Decimal("0.35")

# Items whose price is fixed get no discount. The list comes from the
# FIXED_PRICE_ITEMS environment variable (Railway variable) when it is set:
# one test name per line, or names separated by ";". Otherwise it is the
# "Test" column of fixed_prices.xlsx.
FIXED_PRICES_VAR = "FIXED_PRICE_ITEMS"
FIXED_PRICES_FILE = Path(__file__).parent / "fixed_prices.xlsx"


def normalize(name: str) -> str:
    """Case, repeated spaces and a trailing '*' don't matter when matching names."""
    return " ".join(name.replace("*", " ").split()).lower()


def fixed_items_source() -> str:
    return f"{FIXED_PRICES_VAR} variable" if os.getenv(FIXED_PRICES_VAR, "").strip() else FIXED_PRICES_FILE.name


@cache
def fixed_items() -> frozenset[str]:
    text = os.getenv(FIXED_PRICES_VAR, "")
    if text.strip():
        # Literal "\n" too, in case the newlines were typed as an escape.
        names = re.split(r"\r?\n|;|\\n", text)
    else:
        wb = load_workbook(FIXED_PRICES_FILE, read_only=True, data_only=True)
        try:
            rows = wb.worksheets[0].iter_rows(values_only=True)
            header = [str(h or "").strip().lower() for h in next(rows)]
            col = header.index("test")
            names = [str(r[col]) for r in rows if len(r) > col and r[col]]
        finally:
            wb.close()
    return frozenset(n for n in map(normalize, names) if n)


def is_fixed(desc: str) -> bool:
    return normalize(desc) in fixed_items()


# The mPDF invoices use DejaVu Serif Condensed; the printed ones Times New
# Roman, for which the built-in Times-Bold has the same letter widths.
FONT_DIR = Path(__file__).parent / "fonts"
DEJAVU_BOLD = "DejaVuSerifCondensed-Bold"
pdfmetrics.registerFont(TTFont(DEJAVU_BOLD, str(FONT_DIR / f"{DEJAVU_BOLD}.ttf")))

FIELD = re.compile(r"^(Invoice N\S*|Customer|Sample N\S*|Doctor|Sampling Date|Issued date)\s*:\s*(.*)$")


class InvoiceError(Exception):
    """The PDF looks like an invoice but could not be processed."""


class NotAnInvoice(InvoiceError):
    """The PDF is not an original invoice (lab report, already converted, ...)."""


@dataclass
class Item:
    no: str
    desc: str
    total: Decimal


@dataclass
class Invoice:
    fields: dict[str, str]
    items: list[Item]
    total: Decimal
    label: Run  # "Total:" in the Total row
    value: Run  # the amount in the Total row
    top: Rule  # border above the Total row
    bottom: Rule  # border below the Total row
    left: Rule
    divider: Rule  # between the label cell and the amount cell
    right: Rule
    footer_y: float = 0.0  # a browser page footer below this stays in place on page 2
    page_number: Run | None = None  # the footer's "1/2"


def money(s: str) -> Decimal:
    try:
        return Decimal(s.replace("$", "").replace(",", "").strip())
    except InvalidOperation:
        raise InvoiceError(f"can't read the amount {s!r}") from None


def fmt(d: Decimal) -> str:
    return f"${d:,.2f}"


def cents(d: Decimal) -> Decimal:
    """Half a cent rounds to the even cent, as on the existing discount pages
    (35% of $30.70 is $10.74, of $34.90 is $12.22)."""
    return d.quantize(Decimal("0.01"), rounding=ROUND_HALF_EVEN)


def _lines(runs: list[Run]) -> list[tuple[float, list[Run]]]:
    """Runs grouped by baseline, top to bottom, each line sorted left to right."""
    lines: list[tuple[float, list[Run]]] = []
    for run in sorted(runs, key=lambda r: -r.y):
        if lines and abs(lines[-1][0] - run.y) < 1.0:
            lines[-1][1].append(run)
        else:
            lines.append((run.y, [run]))
    return [(y, sorted(line, key=lambda r: r.x0)) for y, line in lines]


def _join(runs) -> str:
    return " ".join(" ".join(r.text for r in runs).split())


def _verticals_at(rules: list[Rule], y: float) -> list[Rule]:
    """Vertical borders crossing height y, left to right, one per position."""
    found = sorted((r for r in rules if not r.horizontal and r.a - 0.5 <= y <= r.b + 0.5), key=lambda r: r.mid)
    out: list[Rule] = []
    for r in found:
        if not out or r.mid - out[-1].mid > 2:
            out.append(r)
    return out


def parse(runs: list[Run], rules: list[Rule]) -> Invoice:
    text = _join(runs)
    if "INVOICE" not in text or "Invoice N" not in text:
        unreadable = text.count(UNKNOWN)
        raise NotAnInvoice("no INVOICE header" + (f" ({unreadable} unreadable characters)" if unreadable else ""))
    if "Original Amount" in text:
        raise NotAnInvoice("already converted")

    lines = _lines(runs)

    # Header block: "Label : value" lines.
    fields: dict[str, str] = {}
    for _, line in lines:
        m = FIELD.match(_join(line))
        if m:
            key = m.group(1)
            key = "Invoice No" if key.startswith("Invoice") else "Sample No" if key.startswith("Sample") else key
            fields.setdefault(key, m.group(2).strip())

    # Table: rows between the column header and the "Total:" row.
    header_y = next((y for y, line in lines if any(r.text.strip() == "Item#" for r in line)), None)
    total = None
    if header_y is not None:
        total = next(((y, line) for y, line in lines
                      if y < header_y and any(r.text.strip() == "Total:" for r in line)), None)
    if total is None:
        raise InvoiceError("could not find the item table")
    total_y, total_line = total
    label = next(r for r in total_line if r.text.strip() == "Total:")
    amount = [r for r in total_line if r.x0 > label.x0]
    if not amount:
        raise InvoiceError("could not read the invoice total")

    # Columns: Item# | Desc. | Price | Discount | Total, split by the header row's borders.
    borders = _verticals_at(rules, header_y)
    if len(borders) != 6:
        raise InvoiceError(f"expected 5 table columns, found {len(borders) - 1}")
    inner = borders[1:-1]
    items: list[Item] = []
    for y, line in lines:
        if not total_y < y < header_y:
            continue
        cells = ["", "", "", "", ""]
        for r in line:
            col = sum(1 for b in inner if r.x0 > b.mid)
            cells[col] += " " + r.text
        no, desc, _, _, line_total = (" ".join(c.split()) for c in cells)
        if no:
            if not line_total:
                raise InvoiceError(f"item {no} has no total")
            items.append(Item(no, desc, money(line_total)))
        elif desc and items:  # wrapped description line
            items[-1].desc += " " + desc

    # Borders around the Total row.
    x = label.x0 + 1
    across = [r for r in rules if r.horizontal and r.a - 0.5 <= x <= r.b + 0.5]
    above = [r for r in across if r.mid > total_y]
    below = [r for r in across if r.mid < total_y]
    if not above or not below:
        raise InvoiceError("could not find the Total row borders")
    top, bottom = min(above, key=lambda r: r.mid), max(below, key=lambda r: r.mid)
    verticals = _verticals_at(rules, (top.lo + bottom.hi) / 2)
    dividers = [r for r in verticals if label.x0 < r.mid < amount[0].x0]
    if len(verticals) < 3 or not dividers:
        raise InvoiceError("could not find the Total row borders")

    # A browser print has a page footer ("about:blank   1/2") that must not move.
    footer = [r for r in runs if r.y < 40]
    footer_y = max(r.y + r.size for r in footer) + 2 if footer else 0.0
    page_number = next((r for r in footer if r.text == "1/2"), None)

    inv = Invoice(
        fields, items, money(_join(amount)), label, amount[-1], top, bottom,
        verticals[0], dividers[-1], verticals[-1], footer_y, page_number,
    )
    if not items:
        raise InvoiceError("no items found")
    item_sum = sum((i.total for i in items), Decimal(0))
    if item_sum != inv.total:
        raise InvoiceError(f"items add up to {fmt(item_sum)} but invoice total is {fmt(inv.total)}")
    if not fields.get("Customer"):
        raise InvoiceError("no customer name")
    return inv


def discount_for(inv: Invoice) -> Decimal:
    discountable = sum((i.total for i in inv.items if not is_fixed(i.desc)), Decimal(0))
    return cents(discountable * DISCOUNT_RATE)


def _font_for(run: Run) -> str:
    if "DejaVu" in run.font:
        return DEJAVU_BOLD
    if "arial" in run.font.lower() or "helvetica" in run.font.lower():
        return "Helvetica-Bold"
    return "Times-Bold"


def _summary_overlay(width: float, height: float, inv: Invoice, rows: list[tuple[str, str]]) -> bytes:
    font, size = _font_for(inv.label), inv.label.size
    label_right = inv.label.x1 if inv.label.x1 is not None else (
        inv.label.x0 + pdfmetrics.stringWidth(inv.label.text.strip(), font, size))
    # mPDF text has no known right edge: use the label cell's right padding.
    value_right = inv.value.x1 if inv.value.x1 is not None else inv.right.lo - (inv.divider.lo - label_right)
    row_h = inv.top.lo - inv.bottom.lo
    bottom = inv.top.lo - len(rows) * row_h

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))
    c.setFillColorRGB(0, 0, 0)
    for i in range(len(rows) + 1):
        c.rect(inv.left.lo, inv.top.lo - i * row_h, inv.right.hi - inv.left.lo, inv.top.hi - inv.top.lo,
               stroke=0, fill=1)
    for v in (inv.left, inv.divider, inv.right):
        c.rect(v.lo, bottom, v.hi - v.lo, inv.top.hi - bottom, stroke=0, fill=1)
    c.setFont(font, size)
    for i, (label, value) in enumerate(rows):
        y = inv.label.y - i * row_h
        c.drawRightString(label_right, y, label)
        c.drawRightString(value_right, y, value)
    if inv.page_number:  # browser footer: page "1/2" becomes "2/2"
        r = inv.page_number
        right = r.x1 if r.x1 is not None else r.x0 + 2 * r.size
        c.setFillColorRGB(1, 1, 1)
        c.rect(r.x0 - 0.5, r.y - 0.3 * r.size, right - r.x0 + 1, 1.3 * r.size, stroke=0, fill=1)
        c.setFillColorRGB(0, 0, 0)
        c.setFont("Helvetica", r.size)
        c.drawString(r.x0, r.y, "2/2")
    c.showPage()
    c.save()
    return buf.getvalue()


def generate(pdf_bytes: bytes) -> tuple[str, bytes]:
    """Return (output filename, output PDF bytes)."""
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        page = reader.pages[0]
    except Exception as e:  # corrupt / encrypted / not a PDF
        raise NotAnInvoice(str(e)) from e

    # Page 1 first: lab reports stop here without reading their other pages.
    inv = parse(page_runs(page, reader), page_rules(page, reader))
    if any("Original Amount" in _join(page_runs(p, reader)) for p in reader.pages[1:]):
        raise NotAnInvoice("already converted")

    discount = discount_for(inv)
    fixed = [i.desc for i in inv.items if is_fixed(i.desc)]
    log.info(
        "%s: %d items, fixed-price: %s, discount %s",
        inv.fields["Customer"], len(inv.items), ", ".join(fixed) or "none", fmt(discount),
    )
    rows = [
        ("Original Amount:", fmt(inv.total)),
        ("Total for customer:", fmt(inv.total)),
        ("Discount:", fmt(discount)),
        ("Total:", fmt(inv.total - discount)),
    ]
    shift = (len(rows) - 1) * (inv.top.lo - inv.bottom.lo)

    width = float(page.mediabox.width)
    height = float(page.mediabox.height)
    original = page.get_contents().get_data()

    writer = PdfWriter()
    page1 = writer.add_page(page)

    # Page 2 draws the original content again in clipped parts and shares
    # page 1's fonts and images, so the logo and QR code are stored only once.
    resources = page1["/Resources"]
    if not hasattr(resources, "idnum"):
        resources = writer._add_object(resources)
        page1[NameObject("/Resources")] = resources
    page2 = writer.add_blank_page(width, height)
    page2[NameObject("/Resources")] = resources
    big = 10000
    lower = inv.footer_y or -big
    parts = [
        # Everything above the Total row, in place.
        f"q 0 {inv.top.lo:.3f} {width:.3f} {big} re W n\n".encode(), original, b"\nQ\n",
        # Everything below the Total row (down to any page footer), moved down.
        f"q 1 0 0 1 0 {-shift:.3f} cm 0 {lower:.3f} {width:.3f} {inv.bottom.lo - lower:.3f} re W n\n".encode(),
        original, b"\nQ\n",
    ]
    if inv.footer_y:  # the page footer, in place
        parts += [f"q 0 {-big} {width:.3f} {inv.footer_y + big:.3f} re W n\n".encode(), original, b"\nQ\n"]
    stream = DecodedStreamObject()
    stream.set_data(b"".join(parts))
    page2[NameObject("/Contents")] = writer._add_object(stream)

    page2.merge_page(PdfReader(io.BytesIO(_summary_overlay(width, height, inv, rows))).pages[0])

    out = io.BytesIO()
    writer.write(out)

    name = re.sub(r'[\\/:*?"<>|]+', " ", inv.fields["Customer"]).strip()
    return f"{name}.pdf", out.getvalue()


if __name__ == "__main__":
    # Usage: python invoice.py OUT_DIR invoice1.pdf [invoice2.pdf ...]
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    out_dir = Path(sys.argv[1])
    out_dir.mkdir(parents=True, exist_ok=True)
    for src in sys.argv[2:]:
        try:
            name, data = generate(Path(src).read_bytes())
        except InvoiceError as e:
            print(f"{src}: skipped ({type(e).__name__}: {e})")
            continue
        (out_dir / name).write_bytes(data)
        print(f"{src} -> {out_dir / name}")
