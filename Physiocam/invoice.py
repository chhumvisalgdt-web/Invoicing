"""Turn an original CAMMED LAB invoice PDF into the 2-page Physiocam version.

Page 1 is the original page, untouched.
Page 2 is the same page with the "Total:" row replaced by
Original Amount / Total for customer / Discount / Total, and everything
below the table moved down to make room.
"""

import io
import logging
import os
import re
import sys
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from functools import cache
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ContentStream, DecodedStreamObject, NameObject
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

log = logging.getLogger(__name__)

DISCOUNT_RATE = Decimal("0.35")

# Items whose price is fixed get no discount. The list comes from the
# FIXED_PRICE_ITEMS environment variable (Railway variable) when it is set:
# one test name per line, or names separated by ";". Otherwise it is the
# "Test" column of fixed_prices.xlsx.
FIXED_PRICES_VAR = "FIXED_PRICE_ITEMS"
FIXED_PRICES_FILE = Path(__file__).parent / "fixed_prices.xlsx"


def normalize(name: str) -> str:
    """Case, repeated spaces and a trailing '*' don't matter when matching names."""
    return " ".join(name.replace("*", " ").split()).lower()


def fixed_items_source() -> str:
    return f"{FIXED_PRICES_VAR} variable" if os.getenv(FIXED_PRICES_VAR, "").strip() else FIXED_PRICES_FILE.name


@cache
def fixed_items() -> frozenset[str]:
    text = os.getenv(FIXED_PRICES_VAR, "")
    if text.strip():
        # Literal "\n" too, in case the newlines were typed as an escape.
        names = re.split(r"\r?\n|;|\\n", text)
    else:
        wb = load_workbook(FIXED_PRICES_FILE, read_only=True, data_only=True)
        try:
            rows = wb.worksheets[0].iter_rows(values_only=True)
            header = [str(h or "").strip().lower() for h in next(rows)]
            col = header.index("test")
            names = [str(r[col]) for r in rows if len(r) > col and r[col]]
        finally:
            wb.close()
    return frozenset(n for n in map(normalize, names) if n)


def is_fixed(desc: str) -> bool:
    return normalize(desc) in fixed_items()

FONT_DIR = Path(__file__).parent / "fonts"
BOLD = "DejaVuSerifCondensed-Bold"
FONT_SIZE = 11
pdfmetrics.registerFont(TTFont(BOLD, str(FONT_DIR / f"{BOLD}.ttf")))

# Column boundaries of the mPDF invoice table (points from the left edge).
COL_PRICE = 348.589
COL_DISCOUNT = 399.538
COL_TOTAL = 475.962
LINE_WIDTH = 0.567


class InvoiceError(Exception):
    """The PDF looks like an invoice but could not be processed."""


class NotAnInvoice(InvoiceError):
    """The PDF is not an original invoice (lab report, already converted, ...)."""


@dataclass
class Item:
    no: str
    desc: str
    total: Decimal


@dataclass
class Invoice:
    fields: dict[str, str] = field(default_factory=dict)
    items: list[Item] = field(default_factory=list)
    total: Decimal = Decimal(0)
    total_baseline: float = 0.0
    total_label_x: float = 0.0  # where "Total:" starts
    total_value_x: float = 0.0  # where " $xx.xx" starts
    total_value_text: str = ""


def money(s: str) -> Decimal:
    return Decimal(s.replace("$", "").replace(",", "").strip())


def fmt(d: Decimal) -> str:
    return f"${d:,.2f}"


def cents(d: Decimal) -> Decimal:
    return d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _text_runs(page) -> list[tuple[float, float, str]]:
    runs = []

    def visit(text, cm, tm, _fd, _fs):
        if text.strip():
            x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
            y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
            runs.append((x, y, text))

    page.extract_text(visitor_text=visit)
    return runs


def parse(page) -> Invoice:
    runs = _text_runs(page)
    all_text = " ".join(t for _, _, t in runs)
    if "INVOICE" not in all_text or "Invoice N" not in all_text:
        raise NotAnInvoice("no INVOICE header")
    if "Original Amount" in all_text:
        raise NotAnInvoice("already converted")

    inv = Invoice()

    # Header block: "Label" ... " :" ... " value" on the same baseline.
    lines: dict[float, list[tuple[float, str]]] = {}
    for x, y, t in runs:
        lines.setdefault(round(y, 1), []).append((x, t))
    for y in sorted(lines, reverse=True):
        parts = [t for _, t in sorted(lines[y])]
        joined = "".join(parts).strip()
        m = re.match(r"^(Invoice N\S*|Customer|Sample N\S*|Doctor|Sampling Date|Issued date)\s*:\s*(.*)$", joined)
        if m:
            key = m.group(1)
            key = "Invoice No" if key.startswith("Invoice") else "Sample No" if key.startswith("Sample") else key
            inv.fields[key] = m.group(2).strip()

    # Table: rows between the column header and the "Total:" row.
    header_y = next((y for x, y, t in runs if t.strip() == "Item#"), None)
    total_run = next(((x, y) for x, y, t in runs if t.strip() == "Total:"), None)
    if header_y is None or total_run is None:
        raise InvoiceError("could not find the item table")
    inv.total_label_x, inv.total_baseline = total_run

    for y in sorted(lines, reverse=True):
        if not (inv.total_baseline + 1 < y < header_y - 1):
            continue
        cells = sorted(lines[y])
        no = "".join(t for x, t in cells if x < 90).strip()
        desc = "".join(t for x, t in cells if 90 <= x < COL_PRICE).strip()
        line_total = "".join(t for x, t in cells if x >= COL_TOTAL).strip()
        if no:
            if not line_total:
                raise InvoiceError(f"item {no} has no total")
            inv.items.append(Item(no, desc, money(line_total)))
        elif desc and inv.items:  # wrapped description line
            inv.items[-1].desc += " " + desc

    value = [(x, t) for x, y, t in runs if abs(y - inv.total_baseline) < 1 and x > COL_TOTAL]
    if not value:
        raise InvoiceError("could not read the invoice total")
    inv.total_value_x, inv.total_value_text = value[0]
    inv.total = money(inv.total_value_text)

    if not inv.items:
        raise InvoiceError("no items found")
    item_sum = sum((i.total for i in inv.items), Decimal(0))
    if item_sum != inv.total:
        raise InvoiceError(f"items add up to {fmt(item_sum)} but invoice total is {fmt(inv.total)}")
    if not inv.fields.get("Customer"):
        raise InvoiceError("no customer name")
    return inv


def _table_lines(page, reader) -> tuple[list[float], list[float]]:
    """Return y of horizontal and x of vertical line segments drawn on the page."""
    hs, vs = set(), set()
    cur = None
    for operands, op in ContentStream(page.get_contents(), reader).operations:
        if op == b"m":
            cur = tuple(float(v) for v in operands)
        elif op == b"l" and cur:
            x, y = (float(v) for v in operands)
            if abs(y - cur[1]) < 0.01:
                hs.add(round(y, 3))
            elif abs(x - cur[0]) < 0.01:
                vs.add(round(x, 3))
            cur = (x, y)
    return sorted(hs), sorted(vs)


def discount_for(inv: Invoice) -> Decimal:
    discountable = sum((i.total for i in inv.items if not is_fixed(i.desc)), Decimal(0))
    return cents(discountable * DISCOUNT_RATE)


def _summary_overlay(width, height, rows, top, row_h, left, right, label_right, value_right, baseline_off) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))
    c.setLineWidth(LINE_WIDTH)
    c.setStrokeColorRGB(0, 0, 0)
    c.setFont(BOLD, FONT_SIZE)
    for i, (label, value) in enumerate(rows):
        y_top = top - i * row_h
        y_bot = y_top - row_h
        c.line(left, y_top, right, y_top)
        c.line(left, y_bot, right, y_bot)
        c.line(left, y_top, left, y_bot)
        c.line(COL_TOTAL, y_top, COL_TOTAL, y_bot)
        c.line(right, y_top, right, y_bot)
        base = y_top - baseline_off
        c.drawRightString(label_right, base, label)
        c.drawRightString(value_right, base, value)
    c.showPage()
    c.save()
    return buf.getvalue()


def generate(pdf_bytes: bytes) -> tuple[str, bytes]:
    """Return (output filename, output PDF bytes)."""
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        page = reader.pages[0]
    except Exception as e:  # corrupt / encrypted / not a PDF
        raise NotAnInvoice(str(e)) from e
    if any("Original Amount" in (p.extract_text() or "") for p in reader.pages[1:]):
        raise NotAnInvoice("already converted")

    inv = parse(page)

    hs, vs = _table_lines(page, reader)
    above = [y for y in hs if y > inv.total_baseline]
    below = [y for y in hs if y < inv.total_baseline]
    if not above or not below or not vs:
        raise InvoiceError("could not find the table borders")
    row_top, row_bot = min(above), max(below)
    row_h = row_top - row_bot
    left, right = min(vs), max(vs)
    baseline_off = row_top - inv.total_baseline

    label_right = inv.total_label_x + pdfmetrics.stringWidth("Total:", BOLD, FONT_SIZE)
    # Same right padding inside the amount cell as inside the label cell.
    value_right = right - (COL_TOTAL - label_right)

    discount = discount_for(inv)
    fixed = [i.desc for i in inv.items if is_fixed(i.desc)]
    log.info(
        "%s: %d items, fixed-price: %s, discount %s",
        inv.fields["Customer"], len(inv.items), ", ".join(fixed) or "none", fmt(discount),
    )
    rows = [
        ("Original Amount:", fmt(inv.total)),
        ("Total for customer:", fmt(inv.total)),
        ("Discount:", fmt(discount)),
        ("Total:", fmt(inv.total - discount)),
    ]
    shift = (len(rows) - 1) * row_h

    width = float(page.mediabox.width)
    height = float(page.mediabox.height)
    original = page.get_contents().get_data()

    writer = PdfWriter()
    page1 = writer.add_page(page)

    # Page 2 draws the original content twice (clipped) and shares page 1's
    # fonts and images, so the logo and QR code are stored only once.
    resources = page1["/Resources"]
    if not hasattr(resources, "idnum"):
        resources = writer._add_object(resources)
        page1[NameObject("/Resources")] = resources
    page2 = writer.add_blank_page(width, height)
    page2[NameObject("/Resources")] = resources
    big = 10000
    content = b"".join([
        # Everything above the Total row, in place.
        f"q 0 {row_top:.3f} {width:.3f} {big} re W n\n".encode(), original, b"\nQ\n",
        # Everything below the Total row, moved down.
        f"q 1 0 0 1 0 {-shift:.3f} cm 0 {-big} {width:.3f} {row_bot + big:.3f} re W n\n".encode(), original, b"\nQ\n",
    ])
    stream = DecodedStreamObject()
    stream.set_data(content)
    page2[NameObject("/Contents")] = writer._add_object(stream)

    overlay = PdfReader(io.BytesIO(_summary_overlay(
        width, height, rows, row_top, row_h, left, right, label_right, value_right, baseline_off,
    ))).pages[0]
    page2.merge_page(overlay)

    out = io.BytesIO()
    writer.write(out)

    name = re.sub(r'[\\/:*?"<>|]+', " ", inv.fields["Customer"]).strip()
    return f"{name}.pdf", out.getvalue()


if __name__ == "__main__":
    # Usage: python invoice.py OUT_DIR invoice1.pdf [invoice2.pdf ...]
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    out_dir = Path(sys.argv[1])
    out_dir.mkdir(parents=True, exist_ok=True)
    for src in sys.argv[2:]:
        try:
            name, data = generate(Path(src).read_bytes())
        except InvoiceError as e:
            print(f"{src}: skipped ({type(e).__name__}: {e})")
            continue
        (out_dir / name).write_bytes(data)
        print(f"{src} -> {out_dir / name}")
